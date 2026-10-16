"""V41：图像四排从上到下增长；100 nm 下限只约束最下排。"""
import contextlib
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'SEM_0830_0831_compare_V41_bundle'))
import sem_before_after_compare_0831_V41 as sem
from test_v32 import analytic_trenches
np,cv=sem.np,sem.cv2


def increasing_slots(lengths=(21,31,43,55),rim=False,missing=(),dark_columns=(),noise=.2,rim_thickness=3):
    yy,xx=np.indices((400,360))
    background=170.+.03*xx+.02*yy+3.*np.sin(xx/29.)*np.cos(yy/37.)
    signal=np.zeros(xx.shape,np.float32)
    for r,y in enumerate((55,145,235,325)):
        for c,x in enumerate((60,180,300)):
            if (r,c) in missing:continue
            member=np.zeros(xx.shape,np.uint8)
            half=lengths[r]//2
            cv.rectangle(member,(x-8,y-half),(x+8,y+half),1,-1)
            if rim and c not in dark_columns:
                size=2*rim_thickness+1
                outer=cv.dilate(member,cv.getStructuringElement(cv.MORPH_ELLIPSE,(size,size)))
                signal+=65.*(outer-member)
            else:signal-=80.*member
    raw=background+cv.GaussianBlur(signal,(0,0),.8)
    raw+=np.random.default_rng(41).normal(0,noise,xx.shape)
    return np.clip(raw,0,255).astype(np.uint8)


def run_slots(raw,scale=2.,**options):
    pool,_=sem.extract_descriptors_with_rejections(sem.segment_features(raw))
    with sem._v33_detection_context(options),contextlib.redirect_stdout(io.StringIO()):
        return sem._select_and_measure_bottom_slots_v19(raw,pool,scale,'slit.png','single','007','slot210')


def column(lengths=(20.,40.,70.,110.),x=60.,rows=(1,2,3,4)):
    result={}
    for r in rows:
        length=lengths[r-1];y=(55.,145.,235.,325.)[r-1]
        pts=np.array([[x-8,y-length/2],[x,y-length/2],[x+8,y-length/2],
                      [x+8,y],[x+8,y+length/2],[x,y+length/2],
                      [x-8,y+length/2],[x-8,y]],np.float32).reshape(-1,1,2)
        desc=sem.describe_contour(pts)
        desc._v32_edge_audit=dict(edge_support_fraction=1.)
        result[r]=(desc,dict(slot_row=r))
    return result


class LengthRuleTests(unittest.TestCase):
    def audit(self,col,scale=1.):return sem._v41_slot_length_audit(col,np.array([0.,-1.]),scale)

    def test_only_bottom_has_a_100nm_floor_and_upper_rows_can_be_short(self):
        values,issues,audit=self.audit(column())
        self.assertEqual(issues,[])
        np.testing.assert_allclose(list(values.values()),[20.,40.,70.,110.],atol=1e-7)
        self.assertTrue(audit['v41_slit_length_rule_passed'])
        self.assertEqual(audit['v41_slit_minimum_scope'],'bottom_row_only')

    def test_exact_100nm_bottom_is_accepted_and_99_point_9_is_rejected(self):
        for bottom,valid in ((100.,True),(99.9,False)):
            with self.subTest(bottom=bottom):
                _,issues,audit=self.audit(column((20.,40.,70.,bottom)))
                self.assertEqual(not issues,valid)
                self.assertEqual(audit['v41_slit_length_rule_passed'],valid)
                if not valid:self.assertEqual(issues[0]['reason_code'],'slit_bottom_y_length_below_100nm')

    def test_pixel_size_converts_lengths_before_checking_the_floor(self):
        col=column((10.,20.,30.,50.))
        self.assertTrue(self.audit(col,1.)[1])
        self.assertEqual(self.audit(col,2.)[1],[])
        self.assertAlmostEqual(self.audit(col,2.)[0][4],100.,places=7)

    def test_equal_reversed_and_middle_row_inversions_fail_strict_growth(self):
        for lengths in ((20.,20.,70.,110.),(200.,170.,140.,110.),(20.,80.,70.,110.)):
            with self.subTest(lengths=lengths):
                _,issues,_=self.audit(column(lengths))
                self.assertTrue(any(i['reason_code']=='slit_y_length_not_increasing' for i in issues))

    def test_rows_are_kept_in_spatial_order_instead_of_sorted_by_lengths(self):
        col=column();original={r:pair[0].contour.copy() for r,pair in col.items()}
        col[1][0].center[1],col[4][0].center[1]=325.,55.
        self.assertTrue(self.audit(col)[1])
        for r in col:np.testing.assert_array_equal(original[r],col[r][0].contour)

    def test_partial_columns_compare_observed_rows_without_filling_missing_rows(self):
        good=column(rows=(2,4));values,issues,_=self.audit(good)
        self.assertEqual((set(values),issues),({2,4},[]))
        bad=column((20.,130.,70.,110.),rows=(2,4))
        self.assertTrue(self.audit(bad)[1])

    def test_invalid_column_is_removed_before_support_and_other_columns_survive(self):
        cols=[column((20.,40.,70.,90.),60.),column(x=180.),column((25.,50.,80.,120.),300.)]
        rejects=[];kept,_,meta=sem._v41_filter_slot_length_columns(cols,(400,360),1.,rejects)
        self.assertEqual(len(kept),2)
        self.assertEqual(meta['v41_slit_length_floor_rejected_columns'],1)
        self.assertTrue(any(r['reason_code']=='slit_bottom_y_length_below_100nm' for r in rejects))
        self.assertTrue(all(abs(c[4][0].center[0]-60)>20 for c in kept))

    def test_global_row_medians_must_grow_even_when_partial_columns_do(self):
        cols=[column((20.,150.,170.,200.),60.,(1,2,4)),
              column((10.,90.,100.,210.),180.,(1,3,4)),
              column((20.,160.,170.,220.),300.,(2,3,4))]
        self.assertTrue(all(not self.audit(c)[1] for c in cols))
        rejects=[];kept,_,meta=sem._v41_filter_slot_length_columns(cols,(400,360),1.,rejects)
        self.assertEqual(kept,[])
        self.assertEqual(meta['v41_slit_global_length_order_rejected_columns'],3)
        self.assertTrue(any(r['reason_code']=='slit_row_y_lengths_not_increasing' for r in rejects))

    def test_dark_and_bright_rim_scenes_measure_three_growing_bottom_slits(self):
        for rim in (False,True):
            with self.subTest(rim=rim):
                rows,selected,bottom,_,meta=run_slots(increasing_slots(rim=rim))
                self.assertEqual((len(rows),len(selected),len(bottom)),(3,12,3))
                np.testing.assert_allclose([[r['x_width_nm'],r['y_height_nm']] for r in rows],[[34.,110.]]*3,atol=1.)
                self.assertTrue(all(r['v41_slit_length_rule_passed'] for r in rows))
                self.assertEqual(meta['v41_slit_length_rejected_columns'],0)
                if rim:self.assertTrue(all(r['v40_bright_rim_used'] for r in rows))
                for row in rows:
                    measured=json.loads(row['v41_slit_y_lengths_nm_by_row'])
                    self.assertLess(measured['1'],100.)
                    self.assertLess(measured['3'],100.)
                    self.assertGreaterEqual(measured['4'],100.)

    def test_reversed_old_slit_pattern_is_rejected_even_if_bottom_exceeds_100nm(self):
        rows,_,_,rejects,meta=run_slots(increasing_slots((55,43,31,21)),scale=6.)
        self.assertEqual(rows,[])
        self.assertGreater(meta['v41_slit_length_order_rejected_columns'],0)
        self.assertTrue(any(r['reason_code']=='slit_y_length_not_increasing' for r in rejects))

    def test_same_pixels_pass_at_2nm_per_pixel_and_fail_at_1nm_per_pixel(self):
        raw=increasing_slots(rim=True)
        accepted,*_=run_slots(raw,scale=2.)
        rejected,_,_,rejects,_=run_slots(raw,scale=1.)
        self.assertEqual(len(accepted),3)
        self.assertEqual(rejected,[])
        self.assertTrue(any(r['reason_code']=='slit_bottom_y_length_below_100nm' for r in rejects))

    def test_length_hint_and_100_percent_tolerance_cannot_relax_the_floor(self):
        rows,_,_,rejects,_=run_slots(increasing_slots((21,31,43,49)),scale=2.,
                                    slit_length_nm=80.,length_tolerance=1.)
        self.assertEqual(rows,[])
        self.assertTrue(any(r['reason_code']=='slit_bottom_y_length_below_100nm' for r in rejects))

    def test_mixed_dark_and_rim_columns_share_the_same_growth_rule(self):
        rows,selected,_,_,_=run_slots(increasing_slots(rim=True,dark_columns=(0,)))
        self.assertEqual((len(rows),len(selected)),(3,12))
        self.assertEqual(sum(bool(r.get('v40_bright_rim_used')) for r in rows),2)
        self.assertTrue(all(r['v41_slit_length_rule_passed'] for r in rows))

    def test_missing_upper_rows_remain_partial_but_missing_bottom_is_not_measured(self):
        raw=increasing_slots(rim=True,missing=((0,0),(1,1),(2,2),(3,1)))
        rows,_,_,_,_=run_slots(raw)
        self.assertEqual(len(rows),2)
        self.assertTrue(all(r['slot_is_partial_column'] for r in rows))
        self.assertTrue(all(abs(r['center_image_x_px']-180)>20 for r in rows))
        self.assertTrue(all(r['y_height_nm']>=100 for r in rows))
        strict,*_=run_slots(raw,slit_min_column_rows=4)
        self.assertEqual(strict,[])

    def test_missing_entire_bottom_row_is_never_replaced_by_row_three(self):
        rows,*_=run_slots(increasing_slots(rim=True,missing=((3,0),(3,1),(3,2))))
        self.assertEqual(rows,[])

    def test_legacy_bright_polarity_uses_the_same_physical_lengths(self):
        raw=255-increasing_slots()
        with patch.object(sem,'FEATURE_POLARITY','bright'):
            rows,*_=run_slots(raw)
        self.assertEqual(len(rows),3)
        np.testing.assert_allclose([r['y_height_nm'] for r in rows],[110.]*3,atol=1.)

    def test_actual_worker_exports_rule_audit_and_review_flag(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);image=root/'实际 slit.png'
            sem._imwrite_unicode_v19(image,increasing_slots(rim=True))
            image.with_suffix('.txt').write_text('PixelSize=2\n')
            with contextlib.redirect_stdout(io.StringIO()):
                rows,status,_,_,_=sem._v32_process_image_with_timeout(
                    image,'single','007',2,'slot','slot210',root/'ann',timeout_seconds=120.)
            self.assertEqual(len(rows),3)
            self.assertTrue(all(r['v41_slit_length_rule_passed'] for r in rows))
            self.assertTrue(all(r['confidence_flag']=='BRIGHT_RIM_REVIEW' for r in rows))
            self.assertEqual(status['source_algorithm'],'V41_slit_length_order_on_V40_bright_rim')
            self.assertTrue(status['measure_v41_slit_length_rule_enabled'])
            self.assertTrue(Path(status['annotated_path']).is_file())

    def test_zero_result_fallback_never_accepts_short_bottom_slits(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);ann=root/'ann';ann.mkdir()
            for thickness in (3,7):
                with self.subTest(thickness=thickness):
                    image=root/f'short_{thickness}.png'
                    sem._imwrite_unicode_v19(image,increasing_slots((21,31,43,49),rim=True,rim_thickness=thickness))
                    image.with_suffix('.txt').write_text('PixelSize=2\n')
                    with contextlib.redirect_stdout(io.StringIO()):
                        rows,status,rejects=sem._v32_process_v25_assigned(image,'single','007',2,'slot','slot210',ann)
                    self.assertEqual(rows,[])
                    self.assertIn('L1,L2',status['auto_fallback_attempts'])
                    self.assertIn('V39_multiscale_structure',status['auto_fallback_attempts'])
                    self.assertTrue(any(r['reason_code']=='slit_bottom_y_length_below_100nm' for r in rejects))

    def test_real_compare_exports_growing_slits_and_measured_length_delta(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            for stage,lengths in (('before',(21,31,43,55)),('after',(25,35,47,59))):
                folder=root/stage/'007';folder.mkdir(parents=True)
                for i,raw in enumerate((analytic_trenches(),increasing_slots(lengths,rim=True)),1):
                    sem._imwrite_unicode_v19(folder/f'image_{i}.png',raw)
                    (folder/f'image_{i}.txt').write_text('PixelSize=2\n')
            with contextlib.redirect_stdout(io.StringIO()):
                outputs=sem.main_v41(before_root=root/'before',after_root=root/'after',output_root=root/'out',
                                      image_timeout_seconds=120.)
            objects=sem.pd.read_csv(outputs['all_objects'])
            slits=objects[objects.pattern=='slot']
            self.assertEqual(len(slits),6)
            self.assertTrue(slits.v41_slit_length_rule_passed.all())
            pairs=sem.pd.read_csv(outputs['image_pair_comparison'])
            length=pairs[(pairs.pattern=='slot') & (pairs.metric=='slot_y_length_nm')].iloc[0]
            self.assertAlmostEqual(length.delta_after_minus_before_nm,8.,delta=1.)
            settings=json.loads(outputs['settings'].read_text())
            self.assertEqual(settings['v41_slit_length_rule']['minimum_scope'],'bottom_row_only')
            self.assertEqual(settings['v41_slit_length_rule']['bottom_min_y_length_nm'],100.)
            # The inherited compare exporter permits a BOM-only error CSV when
            # no error rows exist; inspect actual error presence in the workbook.
            from openpyxl import load_workbook
            with contextlib.closing(load_workbook(outputs['excel'],read_only=True)) as book:
                self.assertLessEqual(book['errors'].max_row or 0,1)


if __name__=='__main__':unittest.main()
