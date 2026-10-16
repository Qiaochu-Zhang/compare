"""V39 继承的两种运行模式、统计作图及真实 CLI/worker 回归。"""
import contextlib
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'SEM_0830_0831_compare_V39_bundle'))
import sem_before_after_compare_0831_V39 as sem
from test_v30 import gradient_slots
from test_v32 import analytic_trenches, analytic_vias

np, pd = sem.np, sem.pd


class SingleModeTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.input = self.root / '样品 输入'

    def images(self, condition, count=3, root=None, names=None):
        folder = (root or self.input) / condition
        folder.mkdir(parents=True, exist_ok=True)
        for name in reversed(names or [f'scan_{i}.png' for i in range(1, count+1)]):
            (folder / name).touch()
            (folder / name).with_suffix('.txt').write_text('PixelSize=1.0')
        return folder

    @staticmethod
    def worker(**job):
        pattern, condition = job['public_pattern'], job['condition']
        region, stage = job['region'], job['stage']
        # Distinct object counts across regions expose accidental mean-of-means.
        values = [10., 14.] if region == 2 else [30.]
        rows = []
        for index, value in enumerate(values, 1):
            value += stage == 'after'
            rows.append(dict(stage=stage, stage_zh=sem.V32_STAGE_ZH[stage], condition=condition,
                             condition_label=f"{'Sample' if stage == 'single' else 'Condition'} {condition}",
                             region=region, pair_key=f'C{condition}_R{region}_{pattern}', pattern=pattern,
                             pattern_zh=sem.V32_PATTERN_ZH[pattern], image=job['image_path'].name,
                             object_id=index, x_width_nm=value, y_height_nm=value+10,
                             tip_gap_y_nm=value+20, center_image_x_px=index*10., center_image_y_px=20.))
        return rows, {**rows[0], 'status': 'OK'}, [], job['v25_pattern_key'], []

    def run_fake(self, **kwargs):
        with contextlib.redirect_stdout(io.StringIO()), \
                patch.object(sem, '_v32_process_image_with_timeout', side_effect=self.worker), \
                patch.object(sem, '_v35_save_single_plots'), patch.object(sem, '_v32_save_pair_plots'):
            return sem.main_v39(**kwargs)

    def test_inventory_numeric_names_natural_order_and_independent_counts(self):
        expected = ['0', '2', '007', '7', '10', '103']
        for i, condition in enumerate(reversed(expected), 1):
            if condition != '2':
                self.images(condition, i)
        names = ['scan_1.TIF', 'scan_2.PNG', 'scan_3.JPG', 'scan_10.JPEG', 'scan_11.tiff']
        special = self.images('2', names=names)
        (special/'nested').mkdir()
        (special/'nested'/'extra.tif').touch()
        for name in ('notes', '-1', '1.5', '²', '12_suffix'):
            (self.input/name/'99').mkdir(parents=True)
        (self.input/'20').touch()
        inv, errors = sem._v35_build_single_inventory(self.input)
        self.assertTrue(errors.empty)
        self.assertEqual(inv.condition.drop_duplicates().tolist(), expected)
        self.assertEqual(set(inv.stage), {'single'})
        group = inv[inv.condition == '2']
        self.assertEqual(group.image.tolist(), names)
        for _, group in inv.groupby('condition'):
            self.assertEqual(group.sequence_index.tolist(), list(range(1, len(group)+1)))
            self.assertEqual(group.pattern.tolist(), [sem.V32_PATTERN_SEQUENCE[i % 3] for i in range(len(group))])
            self.assertEqual(group.region.tolist(), [2+i//3 for i in range(len(group))])
        self.assertNotEqual(len(inv[inv.condition == '0']), len(inv[inv.condition == '103']))

    def test_empty_sample_and_bad_scale_keep_other_image_positions(self):
        self.images('01', 0)
        folder = self.images('02', 4)
        (folder/'scan_2.txt').unlink()
        (folder/'scan_3.txt').write_text('PixelSize=0')
        inv, errors = sem._v35_build_single_inventory(self.input)
        self.assertEqual(inv.pattern.tolist(), ['trench', 'slot', 'via', 'trench'])
        self.assertEqual(inv.region.tolist(), [2, 2, 2, 3])
        self.assertEqual(errors.error_type.tolist(), ['EMPTY_SAMPLE_FOLDER', 'PIXELSIZE_PREFLIGHT_ERROR', 'PIXELSIZE_PREFLIGHT_ERROR'])

    def test_missing_root_nonnumeric_root_and_all_empty_samples_report_errors(self):
        with self.assertRaisesRegex(FileNotFoundError, 'input 根目录'):
            sem.main_v39(input_root=self.input)
        self.input.mkdir()
        with self.assertRaisesRegex(ValueError, '没有以纯数字命名'):
            sem.main_v39(input_root=self.input)
        self.images('001', 0)
        output = self.root/'empty'
        with contextlib.redirect_stdout(io.StringIO()), self.assertRaisesRegex(RuntimeError, '没有建立任何有效 inventory'):
            sem.main_v39(input_root=self.input, output_root=output)
        self.assertIn('EMPTY_SAMPLE_FOLDER', (output/'processing_errors.csv').read_text(encoding='utf-8-sig'))

    def test_invalid_modes_options_and_nonempty_output_fail_before_measurement(self):
        output = self.root/'out'
        cases = [dict(mode='single'), dict(mode='single', before_root=self.root/'b', input_root=self.input),
                 dict(input_root=self.input, after_root=self.root/'a'), dict(mode='compare', input_root=self.input),
                 dict(mode='invalid'), dict(input_root=self.input, detection_options={'length_tolerance': 0}),
                 dict(input_root=self.input, image_timeout_seconds=301)]
        for kwargs in cases:
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                sem.main_v39(output_root=output, **kwargs)
            self.assertFalse(output.exists())
        self.images('1', 1)
        output.mkdir()
        (output/'keep.txt').write_text('existing user result')
        with self.assertRaisesRegex(ValueError, '输出目录必须不存在或为空'):
            sem.main_v39(input_root=self.input, output_root=output)
        self.assertEqual((output/'keep.txt').read_text(), 'existing user result')

    def test_cli_single_inference_explicit_mode_and_compare_arguments(self):
        for args in (['--input', str(self.input)], ['--mode', 'single', '--input', str(self.input)],
                     ['--before', '/b', '--after', '/a'], ['--mode', 'compare', '--before', '/b', '--after', '/a']):
            with patch.object(sem, 'main_v39') as main:
                self.assertEqual(sem._v32_cli(args), 0)
                self.assertEqual(main.call_args.kwargs['detection_options'].width_tolerance, .5)
                self.assertEqual(main.call_args.kwargs['detection_options'].length_tolerance, 1.)
                if '--input' in args:
                    self.assertEqual(main.call_args.kwargs['input_root'], self.input)
                    self.assertIsNone(main.call_args.kwargs['before_root'])
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(sem._v32_cli(['--mode', 'single']), 1)

    def test_single_exports_do_not_build_comparisons_and_preserve_leading_zero_ids(self):
        for condition, count in [('10', 4), ('007', 6), ('2', 1)]:
            self.images(condition, count)
        with patch.object(sem, '_v32_build_image_pair_comparison', side_effect=AssertionError('comparison called')), \
                patch.object(sem, '_v32_build_overall_comparison', side_effect=AssertionError('comparison called')), \
                patch.object(sem, '_v32_build_object_comparison', side_effect=AssertionError('comparison called')):
            outputs = self.run_fake(input_root=self.input)
        self.assertEqual(outputs['excel'].name, 'SEM_0830_0831_single_V39_results.xlsx')
        self.assertNotIn('image_pair_comparison', outputs)
        for key in ('inventory', 'all_objects', 'image_metrics', 'sample_summary'):
            frame = pd.read_csv(outputs[key], dtype={'condition': str})
            self.assertEqual(frame.condition.drop_duplicates().tolist(), ['2', '007', '10'])
            self.assertFalse(any(c.startswith(('before_', 'after_', 'delta_')) for c in frame.columns))
        summary = pd.read_csv(outputs['sample_summary'], dtype={'condition': str})
        row = summary[(summary.condition == '007') & (summary.metric == 'trench_width_nm')].iloc[0]
        self.assertAlmostEqual(row.mean_nm, 18.)  # (10+14+30)/3, not (12+30)/2
        self.assertAlmostEqual(row.std_nm, np.std([10., 14., 30.], ddof=1))
        self.assertEqual((row.n, row.image_count, row.region_count), (3, 2, 2))
        settings = json.loads(outputs['settings'].read_text())
        self.assertEqual(settings['mode'], 'single')
        self.assertEqual(settings['conditions'], ['2', '007', '10'])
        self.assertNotIn('before_root', settings)
        from openpyxl import load_workbook
        with contextlib.closing(load_workbook(outputs['excel'], read_only=True)) as workbook:
            self.assertIn('sample_summary', workbook.sheetnames)
            self.assertNotIn('object_compare', workbook.sheetnames)
            values = list(workbook['sample_summary'].values)
            index = values[0].index('condition')
            self.assertIn('007', [row[index] for row in values[1:]])

    def test_plots_use_numeric_order_exact_means_and_sample_sd(self):
        objects = []
        for condition in ('10', '007', '2'):
            for region in (2, 3):
                rows, *_ = self.worker(public_pattern='trench', condition=condition, region=region,
                                       stage='single', image_path=Path(f'{region}.png'), v25_pattern_key='trench160')
                objects.extend(rows)
        frame = pd.DataFrame(objects)
        metrics = sem._v32_build_image_metrics(frame)
        summary = sem._v35_build_sample_summary(frame)
        seen = {}
        from matplotlib.figure import Figure
        def capture(fig, path, **kwargs):
            ax = fig.axes[0]
            seen[Path(path).name] = (ax.get_title(), [t.get_text() for t in ax.get_xticklabels()],
                                    [p.get_height() for p in ax.patches],
                                    ax.containers[0].lines[2][0].get_segments(), kwargs['dpi'])
        with patch.object(Figure, 'savefig', autospec=True, side_effect=capture):
            sem._v35_save_single_plots(metrics, summary, self.root/'plots')
        self.assertEqual(len(seen), 10)
        title, labels, bars, segments, dpi = seen['ALL__trench__trench_width_nm__single.png']
        self.assertEqual(labels, ['C2-R2', 'C2-R3', 'C007-R2', 'C007-R3', 'C10-R2', 'C10-R3'])
        np.testing.assert_allclose(bars, [12, 30, 12, 30, 12, 30])
        np.testing.assert_allclose(segments[0][:, 1], [12-np.sqrt(8), 12+np.sqrt(8)])
        np.testing.assert_allclose(segments[1][:, 1], [30, 30])  # n=1: SD remains NaN in data
        self.assertEqual(dpi, 220)
        self.assertIn('Mean +/- SD', title)
        self.assertTrue(metrics[metrics.region == 3].std_nm.isna().all())
        self.assertEqual(seen['SAMPLES__trench__trench_width_nm__single.png'][1], ['C2', 'C007', 'C10'])
        self.assertEqual(seen['SAMPLES__trench__trench_width_nm__single.png'][2], [18., 18., 18.])

    def test_scale_error_timeout_and_corrupt_image_do_not_block_later_images(self):
        folder = self.images('01', 4)
        (folder/'scan_2.txt').unlink()
        seen = []
        def worker(**job):
            seen.append(job['image_path'].name)
            if job['public_pattern'] == 'trench' and job['region'] == 2:
                raise sem.V32ImageTimeout('budget exceeded')
            if job['public_pattern'] == 'via':
                raise RuntimeError('corrupt image')
            return self.worker(**job)
        with contextlib.redirect_stdout(io.StringIO()), \
                patch.object(sem, '_v32_process_image_with_timeout', side_effect=worker), \
                patch.object(sem, '_v35_save_single_plots'):
            outputs = sem.main_v39(input_root=self.input, output_root=self.root/'out')
        self.assertEqual(seen, ['scan_1.png', 'scan_3.png', 'scan_4.png'])
        status = pd.read_csv(outputs['image_status'])
        self.assertEqual(status.status.tolist(), ['SKIPPED_TIMEOUT', 'SKIPPED_PIXELSIZE_PREFLIGHT', 'ERROR', 'OK'])
        data = pd.read_csv(outputs['all_objects'])
        self.assertEqual(data.image.unique().tolist(), ['scan_4.png'])
        errors = pd.read_csv(outputs['errors'])
        self.assertIn('IMAGE_TIMEOUT', errors.error_type.tolist())

    def test_compare_keeps_pair_counts_deltas_and_workbook_tables(self):
        before, after = self.root/'before', self.root/'after'
        self.images('007', 3, before)
        self.images('007', 3, after)
        self.images('10', 1, before)
        self.images('10', 2, after)
        outputs = self.run_fake(before_root=before, after_root=after, output_root=self.root/'compare')
        self.assertNotIn('sample_summary', outputs)
        self.assertEqual(outputs['excel'].name, 'SEM_0830_0831_before_after_V39_results.xlsx')
        pairs = pd.read_csv(outputs['image_pair_comparison'], dtype={'condition': str})
        self.assertEqual(set(pairs.condition), {'007'})
        self.assertEqual(len(pairs), 6)
        np.testing.assert_allclose(pairs.delta_after_minus_before_nm, 1.)
        errors = pd.read_csv(outputs['errors'])
        self.assertEqual(errors.error_type.tolist(), ['IMAGE_COUNT_MISMATCH'])
        from openpyxl import load_workbook
        with contextlib.closing(load_workbook(outputs['excel'], read_only=True)) as workbook:
            self.assertTrue({'image_pair_compare', 'overall_compare', 'object_compare'} <= set(workbook.sheetnames))
        self.assertEqual(json.loads(outputs['settings'].read_text())['mode'], 'compare')

    def test_single_auto_selects_each_image_and_compare_freezes_before_choice(self):
        # Exercise actual auto-selection logic, replacing only detector and process boundary.
        for condition in ('1', '2'):
            self.images(condition, 3)
        calls = []
        def detector(path, stage, condition, region, pattern, key, ann_dir):
            calls.append((stage, condition, pattern, key))
            rows, status, rejects, _, _ = self.worker(image_path=path, stage=stage, condition=condition,
                                                    region=region, public_pattern=pattern, v25_pattern_key=key)
            preferred = 'via40' if condition == '1' else 'via60'
            if pattern == 'via' and key != preferred:
                rows = []
            return rows, status, rejects
        def inline_worker(**kwargs):
            return sem._v32_measure_image_job({**kwargs, 'image_path': str(kwargs['image_path']),
                                               'ann_dir': str(kwargs['ann_dir']),
                                               'detection_options': sem.asdict(kwargs['detection_options'])})
        with contextlib.redirect_stdout(io.StringIO()), \
                patch.object(sem, '_v32_process_v25_assigned', side_effect=detector), \
                patch.object(sem, '_v32_process_image_with_timeout', side_effect=inline_worker), \
                patch.object(sem, '_v35_save_single_plots'), patch.object(sem, '_v32_save_pair_plots'):
            outputs = sem.main_v39(input_root=self.input, output_root=self.root/'single', via_pattern_mode='auto')
            trials = pd.read_csv(outputs['via_model_choice'], dtype={'condition': str})
            self.assertEqual(trials[trials.selected].trial_pattern_key.tolist(), ['via40', 'via60'])
            self.assertEqual(set(trials.stage), {'single'})
            self.assertNotIn('before_image', trials)
            self.assertNotIn('after_image', trials)
            before, after = self.root/'b', self.root/'a'
            self.images('1', 3, before)
            self.images('1', 3, after)
            sem.main_v39(before, after, self.root/'compare', via_pattern_mode='auto')
        self.assertEqual([c[3] for c in calls if c[0] == 'after' and c[2] == 'via'], ['via40'])

    def test_all_scale_errors_still_export_readable_empty_results(self):
        folder = self.images('1', 1)
        (folder/'scan_1.txt').unlink()
        outputs = self.run_fake(input_root=self.input)
        for key in ('all_objects', 'image_metrics', 'sample_summary'):
            self.assertTrue(pd.read_csv(outputs[key]).empty)
        self.assertEqual(pd.read_csv(outputs['image_status']).status.tolist(), ['SKIPPED_PIXELSIZE_PREFLIGHT'])

    def test_real_cli_single_all_patterns_auto_calibration_plots_and_excel(self):
        first = self.images('007', 3)
        second = self.images('10', 1)
        for folder, scale, images in ((first, 2., [analytic_trenches(), gradient_slots(40.), analytic_vias()[0]]),
                                      (second, 1., [analytic_trenches()])):
            for i, gray in enumerate(images, 1):
                sem._imwrite_unicode_v19(folder/f'scan_{i}.png', gray)
                (folder/f'scan_{i}.txt').write_text(f'PixelSize={scale}')
        output = self.root/'结果'
        env = {**os.environ, 'MPLBACKEND': 'Agg', 'OPENBLAS_NUM_THREADS': '1', 'PYTHONDONTWRITEBYTECODE': '1'}
        result = subprocess.run([sys.executable, str(Path(sem.__file__)), '--input', str(self.input),
                                 '--output', str(output), '--via-pattern', 'auto', '--image-timeout-seconds', '60',
                                 '--width-tolerance', '.25', '--length-tolerance', '.8'],
                                capture_output=True, text=True, env=env, timeout=180)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn('V39 | SEM single', result.stdout)
        objects = pd.read_csv(output/'all_object_measurements.csv', dtype={'condition': str})
        self.assertEqual(set(objects.stage), {'single'})
        self.assertTrue((objects.input_width_tolerance == .25).all())
        self.assertTrue((objects.input_length_tolerance == .8).all())
        expected = {('007', 'trench'): (9, 40., 'tip_gap_y_nm', 120.),
                    ('007', 'slot'): (3, 34., 'y_height_nm', 42.),
                    ('007', 'via'): (20, 80., 'y_height_nm', 80.),
                    ('10', 'trench'): (9, 20., 'tip_gap_y_nm', 60.)}
        for (condition, pattern), (count, width, second_metric, height) in expected.items():
            group = objects[(objects.condition == condition) & (objects.pattern == pattern)]
            self.assertEqual(len(group), count, (condition, pattern, result.stdout))
            np.testing.assert_allclose(group.x_width_nm, width, atol=2.)
            np.testing.assert_allclose(group[second_metric], height, atol=2.)
        status = pd.read_csv(output/'image_status.csv')
        self.assertEqual(len(status), 4)
        self.assertTrue(all(Path(p).is_file() for p in status.annotated_path))
        self.assertEqual(len(list((output/'plots').glob('*.png'))), 20)
        self.assertTrue(all(p.stat().st_size > 1000 for p in (output/'plots').glob('*.png')))
        self.assertFalse(list(output.glob('*comparison*')))
        self.assertTrue(pd.read_csv(output/'processing_errors.csv').empty)
        settings = json.loads((output/'settings.json').read_text())
        self.assertEqual(settings['accepted_object_count'], 41)
        self.assertIn('-V39-', settings['script_version'])
        self.assertTrue(settings['v39_zero_recovery_enabled'])
        from openpyxl import load_workbook
        with contextlib.closing(load_workbook(output/'SEM_0830_0831_single_V39_results.xlsx', read_only=True)) as workbook:
            self.assertEqual(workbook['all_objects'].max_row, 42)
            self.assertEqual(workbook['sample_summary'].max_row, 9)

    def test_real_worker_blank_image_exports_no_zero_measurements(self):
        folder = self.images('1', 1)
        sem._imwrite_unicode_v19(folder/'scan_1.png', np.full((200, 300), 150, np.uint8))
        with contextlib.redirect_stdout(io.StringIO()):
            outputs = sem.main_v39(input_root=self.input, output_root=self.root/'blank', image_timeout_seconds=30.)
        self.assertTrue(pd.read_csv(outputs['sample_summary']).empty)
        self.assertTrue(pd.read_csv(outputs['all_objects']).empty)
        self.assertEqual(list((self.root/'blank'/'plots').iterdir()), [])

    def test_real_worker_timeout_does_not_publish_partial_annotations(self):
        folder = self.images('1', 1)
        sem._imwrite_unicode_v19(folder/'scan_1.png', analytic_trenches())
        with self.assertRaises(sem.V32ImageTimeout):
            sem._v32_process_image_with_timeout(folder/'scan_1.png', 'single', '1', 2,
                                                'trench', 'trench160', self.root/'ann', .001)
        self.assertFalse((self.root/'ann').exists())


if __name__ == '__main__':
    unittest.main()
